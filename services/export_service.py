"""
CSV and Excel Data Export Service
"""

import csv
import openpyxl
from pathlib import Path
from config import EXPORTS_DIR

def export_to_csv(filename, headers, rows):
    file_path = EXPORTS_DIR / filename
    with open(file_path, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    return str(file_path)

def export_to_excel(filename, sheet_name, headers, rows):
    file_path = EXPORTS_DIR / filename
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    # Header Row Style
    ws.append(headers)
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        cell.fill = openpyxl.styles.PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    for row in rows:
        ws.append(row)

    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(file_path)
    return str(file_path)
